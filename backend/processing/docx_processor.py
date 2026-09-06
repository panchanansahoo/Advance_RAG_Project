"""
Processor for DOCX (Microsoft Word) files.

Phase 3 component. Extracts paragraphs and tables from Word documents.
"""

from __future__ import annotations

import logging
from typing import Any, Dict, List, Optional
import docx

from backend.processing.base import BaseProcessor, ProcessedContent
from backend.schemas.common import ContentType

logger = logging.getLogger(__name__)


class DOCXProcessor(BaseProcessor):
    """
    Extracts text from DOCX files.
    """

    def supported_extensions(self) -> List[str]:
        return [".docx"]

    async def process(
        self, file_path: str, metadata: Optional[Dict[str, Any]] = None
    ) -> List[ProcessedContent]:
        """
        Process a DOCX file.
        """
        metadata = metadata or {}
        contents: List[ProcessedContent] = []

        try:
            doc = docx.Document(file_path)
        except Exception as e:
            logger.error("Failed to open DOCX file: %s", e)
            raise ValueError(f"Invalid DOCX file: {e}")

        # Extract text from paragraphs
        text_elements = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                text_elements.append(text)

        # Extract text from tables
        for table in doc.tables:
            text_elements.append("\n[Table]")
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                # Avoid totally empty rows
                if any(row_data):
                    text_elements.append(" | ".join(row_data))
            text_elements.append("\n")

        full_text = "\n\n".join(text_elements)

        if full_text.strip():
            contents.append(
                ProcessedContent(
                    text=full_text,
                    content_type=ContentType.TEXT,
                    page_number=1,
                    metadata={
                        **metadata,
                        "source": "python-docx",
                    },
                )
            )

        logger.info("DOCX processed: %s → %d content units", file_path, len(contents))
        return contents
